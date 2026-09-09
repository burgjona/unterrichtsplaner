"""Notenmodul, Meilenstein 1: Rechenkern (src/lib/noten.py) und API (src/routers/noten.py)."""
import pytest

from src.lib import noten as calc


# ---------- Rechenkern (ohne DB) ----------

@pytest.mark.parametrize("text,expected", [
    ("1", 1.0), ("1+", 0.75), ("1-", 1.25),
    ("2+", 1.75), ("2", 2.0), ("2-", 2.25),
    ("6", 6.0), ("6+", 5.75),
    ("2,5", 2.5), ("2.5", 2.5), (" 3 ", 3.0),
    ("", None), (None, None), (3, 3.0),
])
def test_parse_note(text, expected):
    assert calc.parse_note(text) == expected


@pytest.mark.parametrize("text", ["0", "7", "6-", "abc", "+", "1++"])
def test_parse_note_ungueltig(text):
    with pytest.raises(ValueError):
        calc.parse_note(text)


@pytest.mark.parametrize("value,expected", [
    (0.75, "1+"), (1.0, "1"), (1.25, "1-"), (2.25, "2-"), (6.0, "6"),
    (2.4, "2,4"), (2.33, "2,33"), (None, ""),
])
def test_format_note(value, expected):
    assert calc.format_note(value) == expected


def test_average_ignoriert_leere():
    assert calc.average([1.0, None, 3.0]) == 2.0
    assert calc.average([None, None]) is None


def test_gesamt_average_mischt_nach_anteil():
    # 60 % groß (Ø 2,0) + 40 % klein (Ø 3,0) = 2,4
    assert calc.gesamt_average(2.0, 3.0, 60) == 2.4
    assert calc.gesamt_average(2.0, 3.0, 50) == 2.5


def test_gesamt_average_leerer_topf_zaehlt_der_andere():
    assert calc.gesamt_average(None, 3.0, 70) == 3.0
    assert calc.gesamt_average(2.0, None, 70) == 2.0
    assert calc.gesamt_average(None, None) is None


def test_suggest_term_grade_kaufmaennisch():
    assert calc.suggest_term_grade(2.4) == 2.0
    assert calc.suggest_term_grade(2.5) == 3.0
    assert calc.suggest_term_grade(None) is None


# ---------- API ----------

@pytest.fixture
def klasse(client, auth):
    r = client.post("/api/classes", json={
        "name": "8a", "subject": "Deutsch", "grade": 8, "track": "gemischt",
    })
    assert r.status_code == 201, r.text
    cid = r.json()["id"]
    r = client.post(f"/api/classes/{cid}/students/bulk",
                    json={"names": ["Änne Müller", "Bert Straßer"]})
    assert r.status_code == 201, r.text
    return cid, [s["id"] for s in r.json()]


def _item(client, cid, **kw):
    body = {"title": "Balladen", "kind": "Klassenarbeit", "size": "gross",
            "date": "2025-11-14"}
    body.update(kw)
    r = client.post(f"/api/classes/{cid}/grade-items", json=body)
    assert r.status_code == 201, r.text
    return r.json()


def test_leistung_anlegen_halbjahr_aus_datum(client, klasse):
    cid, _ = klasse
    assert _item(client, cid, date="2025-11-14")["term"] == "HJ1"
    assert _item(client, cid, date="2026-01-20")["term"] == "HJ1"
    assert _item(client, cid, date="2026-03-05")["term"] == "HJ2"
    # ausdrücklich gesetztes Halbjahr schlägt die Vorbelegung
    assert _item(client, cid, date="2026-03-05", term="HJ1")["term"] == "HJ1"


def test_unbekannter_anlass_wird_abgewiesen(client, klasse):
    cid, _ = klasse
    r = client.post(f"/api/classes/{cid}/grade-items", json={
        "title": "x", "kind": "Wandertag", "size": "klein", "date": "2025-11-14"})
    assert r.status_code == 422


def test_note_mit_tendenz_speichern_und_anzeigen(client, klasse):
    cid, students = klasse
    item = _item(client, cid)
    r = client.put(f"/api/grade-items/{item['id']}/students/{students[0]}/grade",
                   json={"value": "2-", "comment": "Aufsatz sehr strukturiert"})
    assert r.status_code == 200, r.text
    assert r.json()["value"] == 2.25
    assert r.json()["label"] == "2-"
    assert r.json()["comment"] == "Aufsatz sehr strukturiert"


def test_leere_note_loescht_die_zelle(client, klasse):
    cid, students = klasse
    item = _item(client, cid)
    client.put(f"/api/grade-items/{item['id']}/students/{students[0]}/grade",
               json={"value": "3"})
    r = client.put(f"/api/grade-items/{item['id']}/students/{students[0]}/grade",
                   json={"value": None})
    assert r.status_code == 200 and r.json() is None
    m = client.get(f"/api/classes/{cid}/noten", params={"term": "HJ1"}).json()
    assert m["grades"] == []


def test_note_fremder_schueler_wird_abgewiesen(client, klasse, auth):
    cid, _ = klasse
    r = client.post("/api/classes", json={"name": "9b", "subject": "WTH", "grade": 9})
    other = r.json()["id"]
    r = client.post(f"/api/classes/{other}/students", json={"name": "Fremd"})
    fremd = r.json()["id"]
    item = _item(client, cid)
    r = client.put(f"/api/grade-items/{item['id']}/students/{fremd}/grade", json={"value": "1"})
    assert r.status_code == 404


def test_matrix_durchschnitte_und_gewichtung(client, klasse):
    cid, students = klasse
    gross = _item(client, cid, size="gross", title="Klassenarbeit")
    klein = _item(client, cid, size="klein", title="Hausaufgabe", kind="Hausaufgabe")
    s = students[0]
    client.put(f"/api/grade-items/{gross['id']}/students/{s}/grade", json={"value": "2"})
    client.put(f"/api/grade-items/{klein['id']}/students/{s}/grade", json={"value": "3"})

    m = client.get(f"/api/classes/{cid}/noten", params={"term": "HJ1"}).json()
    assert m["grossAnteil"] == 50
    summary = next(x for x in m["summaries"] if x["studentId"] == s)
    assert summary["avgGross"] == 2.0
    assert summary["avgKlein"] == 3.0
    assert summary["avgGesamt"] == 2.5
    assert summary["suggestedTermGrade"] == 3.0

    r = client.put(f"/api/classes/{cid}/noten-gewichtung", json={"grossAnteil": 70})
    assert r.status_code == 200 and r.json()["grossAnteil"] == 70
    m = client.get(f"/api/classes/{cid}/noten", params={"term": "HJ1"}).json()
    summary = next(x for x in m["summaries"] if x["studentId"] == s)
    assert summary["avgGesamt"] == 2.3          # 2*0,7 + 3*0,3
    assert summary["suggestedTermGrade"] == 2.0


def test_matrix_filtert_nach_halbjahr(client, klasse):
    cid, students = klasse
    hj1 = _item(client, cid, date="2025-11-14")
    hj2 = _item(client, cid, date="2026-03-05")
    client.put(f"/api/grade-items/{hj1['id']}/students/{students[0]}/grade", json={"value": "1"})
    client.put(f"/api/grade-items/{hj2['id']}/students/{students[0]}/grade", json={"value": "5"})

    m1 = client.get(f"/api/classes/{cid}/noten", params={"term": "HJ1"}).json()
    m2 = client.get(f"/api/classes/{cid}/noten", params={"term": "HJ2"}).json()
    assert [i["id"] for i in m1["items"]] == [hj1["id"]]
    assert next(x for x in m1["summaries"] if x["studentId"] == students[0])["avgGesamt"] == 1.0
    assert next(x for x in m2["summaries"] if x["studentId"] == students[0])["avgGesamt"] == 5.0


def test_schnellerfassung_bulk(client, klasse):
    cid, students = klasse
    item = _item(client, cid)
    r = client.put(f"/api/grade-items/{item['id']}/grades", json={"grades": [
        {"studentId": students[0], "value": "1-"},
        {"studentId": students[1], "value": "4"},
    ]})
    assert r.status_code == 200, r.text
    assert sorted(g["label"] for g in r.json()) == ["1-", "4"]


def test_zeugnisnote_setzen_und_zuruecknehmen(client, klasse):
    cid, students = klasse
    item = _item(client, cid)
    s = students[0]
    client.put(f"/api/grade-items/{item['id']}/students/{s}/grade", json={"value": "3"})

    r = client.put(f"/api/classes/{cid}/students/{s}/term-grade", params={"term": "HJ1"},
                   json={"value": "2", "comment": "deutliche Steigerung im Halbjahr"})
    assert r.status_code == 200, r.text
    assert r.json()["value"] == 2.0

    m = client.get(f"/api/classes/{cid}/noten", params={"term": "HJ1"}).json()
    summary = next(x for x in m["summaries"] if x["studentId"] == s)
    assert summary["termGrade"] == 2.0
    assert summary["suggestedTermGrade"] == 3.0   # Vorschlag bleibt sichtbar daneben

    r = client.put(f"/api/classes/{cid}/students/{s}/term-grade", params={"term": "HJ1"},
                   json={"value": None})
    assert r.status_code == 200 and r.json() is None
    m = client.get(f"/api/classes/{cid}/noten", params={"term": "HJ1"}).json()
    assert next(x for x in m["summaries"] if x["studentId"] == s)["termGrade"] is None


def test_leistung_loeschen_nimmt_noten_mit(client, klasse):
    cid, students = klasse
    item = _item(client, cid)
    client.put(f"/api/grade-items/{item['id']}/students/{students[0]}/grade", json={"value": "2"})
    assert client.delete(f"/api/grade-items/{item['id']}").status_code == 204
    m = client.get(f"/api/classes/{cid}/noten", params={"term": "HJ1"}).json()
    assert m["items"] == [] and m["grades"] == []


def test_umlaute_bleiben_erhalten(client, klasse):
    cid, students = klasse
    item = _item(client, cid, title="Kurzprüfung: Groß- und Kleinschreibung",
                 kind="Leistungskontrolle", note="Nachschreiber Straßer")
    assert item["title"] == "Kurzprüfung: Groß- und Kleinschreibung"
    assert item["note"] == "Nachschreiber Straßer"
    m = client.get(f"/api/classes/{cid}/noten", params={"term": "HJ1"}).json()
    assert m["students"][0]["name"] == "Änne Müller"


def test_noten_sind_sync_faehig(client, klasse):
    cid, students = klasse
    item = _item(client, cid)
    client.put(f"/api/grade-items/{item['id']}/students/{students[0]}/grade", json={"value": "2"})
    r = client.get("/api/sync/changes",
                   params={"since": 0, "entities": "grade_items,grades,term_grades"})
    assert r.status_code == 200, r.text
    types = {c["entityType"] for c in r.json()["changes"]}
    assert "grade_items" in types and "grades" in types


# ---------- Word-Export (Meilenstein 3) ----------

DOCX_MEDIA = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _docx_text(payload: bytes) -> str:
    """Sichtbarer Text eines .docx – reicht, um Inhalte zu prüfen, ohne das Layout zu fixieren."""
    from io import BytesIO

    from docx import Document

    doc = Document(BytesIO(payload))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


@pytest.fixture
def noten_daten(client, klasse):
    """Klasse mit zwei Leistungen (groß + klein) und Noten für den ersten Schüler."""
    cid, students = klasse
    gross = _item(client, cid, size="gross", title="Klassenarbeit Balladen")
    klein = _item(client, cid, size="klein", title="Vortrag Erlkönig", kind="Referat")
    client.put(f"/api/grade-items/{gross['id']}/students/{students[0]}/grade",
               json={"value": "2", "comment": "sicherer Aufbau"})
    client.put(f"/api/grade-items/{klein['id']}/students/{students[0]}/grade",
               json={"value": "3"})
    client.put(f"/api/grade-items/{gross['id']}/students/{students[1]}/grade",
               json={"value": "4"})
    return cid, students, gross, klein


def test_export_matrix(client, noten_daten):
    cid, students, gross, klein = noten_daten
    r = client.get(f"/api/classes/{cid}/noten/export/matrix", params={"term": "HJ1"})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == DOCX_MEDIA
    text = _docx_text(r.content)
    assert "Notenübersicht" in text
    assert "Klassenarbeit Balladen" in text and "Vortrag Erlkönig" in text
    assert "14.11.2025" in text                    # deutsches Datum, nicht ISO
    assert "2025-11-14" not in text
    assert "Änne Müller" in text                   # Umlaute unverändert
    assert "2,50" in text                          # Ø gesamt des ersten Schülers


def test_export_einzelblatt_je_schueler(client, noten_daten):
    cid, students, gross, klein = noten_daten
    r = client.get(f"/api/classes/{cid}/noten/export/schueler", params={"term": "HJ1"})
    assert r.status_code == 200, r.text
    text = _docx_text(r.content)
    assert "Änne Müller" in text and "Bert Straßer" in text
    assert "sicherer Aufbau" in text               # Hinweisfeld der Note
    # gross 2 + klein 3 bei 50:50 ergibt genau 2,50 -- ein Grenzfall, deshalb beide Noten.
    assert "Rechnerischer Vorschlag für die Zeugnisnote: 2 oder 3" in text


def test_export_auswertung_einer_leistung(client, noten_daten):
    cid, students, gross, klein = noten_daten
    r = client.get(f"/api/grade-items/{gross['id']}/export")
    assert r.status_code == 200, r.text
    text = _docx_text(r.content)
    assert "Klassenarbeit Balladen" in text
    assert "Notenspiegel" in text
    assert "Bewertet: 2 von 2" in text
    assert "Durchschnitt: 3,00" in text            # Noten 2 und 4


def test_export_zeugnisnotenliste(client, noten_daten):
    cid, students, gross, klein = noten_daten
    client.put(f"/api/classes/{cid}/students/{students[0]}/term-grade",
               params={"term": "HJ1"}, json={"value": "2", "comment": "deutliche Steigerung"})
    r = client.get(f"/api/classes/{cid}/noten/export/zeugnis", params={"term": "HJ1"})
    assert r.status_code == 200, r.text
    text = _docx_text(r.content)
    assert "Zeugnisnoten" in text
    assert "deutliche Steigerung" in text
    assert "Vorschlag" in text


def test_export_dateiname_mit_umlaut(client, noten_daten):
    cid, _, _, _ = noten_daten
    r = client.get(f"/api/classes/{cid}/noten/export/matrix", params={"term": "HJ1"})
    cd = r.headers["content-disposition"]
    assert "filename*=UTF-8''" in cd
    assert "Noten%C3%BCbersicht" in cd             # RFC 5987 trägt das ü
    assert "Noten_bersicht" in cd                  # ASCII-Fallback daneben


def test_export_leeres_halbjahr_bleibt_erzeugbar(client, klasse):
    cid, _ = klasse
    r = client.get(f"/api/classes/{cid}/noten/export/matrix", params={"term": "HJ2"})
    assert r.status_code == 200, r.text
    assert "noch keine Leistung erfasst" in _docx_text(r.content)


def test_export_fremde_klasse_404(client, auth):
    assert client.get("/api/classes/999/noten/export/matrix").status_code == 404
    assert client.get("/api/grade-items/999/export").status_code == 404



def test_matrix_export_ist_querformat(client, noten_daten):
    """Die Übersicht braucht Querformat – Seitenmaße UND w:orient.

    Regression: python-docx kennt section.orientation, nicht section.orient. Ein Tippfehler
    dort legt still ein neues Attribut an, statt die Ausrichtung zu setzen.
    """
    from io import BytesIO

    from docx import Document
    from docx.enum.section import WD_ORIENT

    cid, _, _, _ = noten_daten
    r = client.get(f"/api/classes/{cid}/noten/export/matrix", params={"term": "HJ1"})
    section = Document(BytesIO(r.content)).sections[0]
    assert section.orientation == WD_ORIENT.LANDSCAPE
    assert section.page_width > section.page_height


def test_einzelblatt_hat_eine_seite_je_schueler(client, noten_daten):
    from io import BytesIO

    from docx import Document

    cid, students, _, _ = noten_daten
    r = client.get(f"/api/classes/{cid}/noten/export/schueler", params={"term": "HJ1"})
    doc = Document(BytesIO(r.content))
    breaks = sum('type="page"' in run._element.xml
                 for p in doc.paragraphs for run in p.runs)
    assert breaks == len(students) - 1      # Umbruch zwischen den Blättern, nicht davor


# ---------- Grenzfälle: Ø genau zwischen zwei Ganznoten ----------

@pytest.mark.parametrize("gesamt,grenzfall", [
    (2.4, False), (2.49, False), (2.5, True), (2.51, False),
    (1.5, True), (5.5, True), (6.0, False), (0.75, False), (None, False),
])
def test_is_borderline(gesamt, grenzfall):
    assert calc.is_borderline(gesamt) is grenzfall


def test_summarize_meldet_grenzfall(client, klasse):
    """Eine große 2 und eine kleine 3 bei 50:50 ergeben genau 2,50."""
    cid, students = klasse
    gross = _item(client, cid, size="gross")
    klein = _item(client, cid, size="klein", kind="Hausaufgabe")
    s = students[0]
    client.put(f"/api/grade-items/{gross['id']}/students/{s}/grade", json={"value": "2"})
    client.put(f"/api/grade-items/{klein['id']}/students/{s}/grade", json={"value": "3"})

    m = client.get(f"/api/classes/{cid}/noten", params={"term": "HJ1"}).json()
    summary = next(x for x in m["summaries"] if x["studentId"] == s)
    assert summary["avgGesamt"] == 2.5
    assert summary["termGradeBorderline"] is True
    # Der Vorschlag bleibt die obere Note – die Oberfläche macht daraus "2 oder 3".
    assert summary["suggestedTermGrade"] == 3.0

    # Ohne Grenzlage keine Markierung.
    client.put(f"/api/grade-items/{klein['id']}/students/{s}/grade", json={"value": "3-"})
    m = client.get(f"/api/classes/{cid}/noten", params={"term": "HJ1"}).json()
    assert next(x for x in m["summaries"] if x["studentId"] == s)["termGradeBorderline"] is False


def test_export_zeigt_beide_moeglichkeiten(client, klasse):
    cid, students = klasse
    gross = _item(client, cid, size="gross")
    klein = _item(client, cid, size="klein", kind="Hausaufgabe")
    s = students[0]
    client.put(f"/api/grade-items/{gross['id']}/students/{s}/grade", json={"value": "2"})
    client.put(f"/api/grade-items/{klein['id']}/students/{s}/grade", json={"value": "3"})

    text = _docx_text(client.get(f"/api/classes/{cid}/noten/export/zeugnis",
                                 params={"term": "HJ1"}).content)
    assert "2 oder 3" in text
    text = _docx_text(client.get(f"/api/classes/{cid}/noten/export/schueler",
                                 params={"term": "HJ1"}).content)
    assert "Rechnerischer Vorschlag für die Zeugnisnote: 2 oder 3" in text


def test_notenspiegel_weist_halbwerte_gesondert_aus(client, klasse):
    """2,5 landet in keinem Balken; die Anteile beziehen sich auf die übrigen Noten."""
    cid, students = klasse
    item = _item(client, cid)
    client.put(f"/api/grade-items/{item['id']}/students/{students[0]}/grade", json={"value": "2,5"})
    client.put(f"/api/grade-items/{item['id']}/students/{students[1]}/grade", json={"value": "3"})

    text = _docx_text(client.get(f"/api/grade-items/{item['id']}/export").content)
    zeilen = text.split("\n")
    idx = zeilen.index("Anzahl")
    assert zeilen[idx:idx + 7] == ["Anzahl", "0", "0", "1", "0", "0", "0"]   # nur die 3
    anteil = zeilen.index("Anteil")
    assert zeilen[anteil:anteil + 7] == ["Anteil", "0 %", "0 %", "100 %", "0 %", "0 %", "0 %"]
    assert "Auf der Grenze: 1 (2,50)" in text
    assert "Bewertet: 2 von 2" in text        # der Halbwert zählt weiter als bewertet
    assert "Durchschnitt: 2,75" in text       # und geht in den Durchschnitt ein


def test_notenspiegel_ohne_halbwerte_ohne_zusatzzeile(client, klasse):
    cid, students = klasse
    item = _item(client, cid)
    client.put(f"/api/grade-items/{item['id']}/students/{students[0]}/grade", json={"value": "2"})
    text = _docx_text(client.get(f"/api/grade-items/{item['id']}/export").content)
    assert "Auf der Grenze" not in text
