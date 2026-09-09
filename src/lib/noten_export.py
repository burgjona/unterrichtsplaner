"""Word-Export des Notenmoduls (docs/konzept_noten.md, Abschnitt 7).

Vier Vorlagen, alle auf python-docx (bereits Dependency, vgl. asuv_export.py):

1. build_matrix_docx        Notenübersicht der Klasse, Querformat: Schüler × Leistungen
2. build_student_sheets_docx Einzelblatt je Schüler, eine Seite pro Kind
3. build_item_docx          Auswertung einer Leistung mit Notenspiegel
4. build_term_docx          Zeugnisnoten-Liste, kompakt

Alle vier bekommen dasselbe `ctx`-Dict aus src/routers/noten.py (siehe collect_export_data)
und liefern die Datei als bytes. Datumsangaben durchgängig TT.MM.JJJJ, Umlaute bleiben
erhalten – auch in den Dateinamen, dafür sorgt der RFC-5987-Header im Router.
"""
from datetime import date
from io import BytesIO
from typing import List, Optional

from . import noten as calc

SIZE_LABEL = {"gross": "groß", "klein": "klein"}
TERM_LABEL = {"HJ1": "1. Halbjahr", "HJ2": "2. Halbjahr"}


def _de(iso: Optional[str]) -> str:
    """'YYYY-MM-DD' → 'TT.MM.JJJJ'."""
    if not iso or len(iso) < 10:
        return ""
    return f"{iso[8:10]}.{iso[5:7]}.{iso[0:4]}"


def _num(value: Optional[float]) -> str:
    """Durchschnitt als deutsche Dezimalzahl; '–', solange nichts bewertet ist."""
    return "–" if value is None else f"{value:.2f}".replace(".", ",")


def _note(value: Optional[float]) -> str:
    return calc.format_note(value) or "–"


def _suggestion(summary: dict) -> str:
    """Zeugnisnoten-Vorschlag; bei einem Grenzfall beide Möglichkeiten statt einer Zahl.

    Ø 2,50 heißt nicht „3“, sondern „hier musst du dich entscheiden“ – das soll auch im
    Ausdruck stehen, nicht nur auf dem Bildschirm.
    """
    value = summary.get("suggested_term_grade")
    if value is None:
        return "–"
    if summary.get("term_grade_borderline"):
        return f"{int(value) - 1} oder {int(value)}"
    return _note(value)


def _doc(landscape: bool = False):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)

    sec = doc.sections[0]
    if landscape:
        # Beides nötig: die Seitenmaße tauschen macht das Blatt quer, orientation setzt
        # zusätzlich das w:orient-Attribut, an dem Word und der Drucker sich orientieren.
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, attr, Cm(1.5))
    return doc


def _title(doc, text: str, sub: str = ""):
    from docx.shared import Pt

    p = doc.add_paragraph(text)
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(15)
    if sub:
        s = doc.add_paragraph(sub)
        s.runs[0].font.size = Pt(9)


def _table(doc, header: List[str], rows: List[List[str]], widths_cm=None):
    """Tabelle mit fetter Kopfzeile im Standard-Gitter."""
    from docx.shared import Cm, Pt

    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(text))
            run.font.size = Pt(9)
    if widths_cm:
        # python-docx setzt Spaltenbreiten nur zuverlässig, wenn sie an JEDER Zelle stehen.
        for row in table.rows:
            for cell, cm in zip(row.cells, widths_cm):
                if cm:
                    cell.width = Cm(cm)
    return table


def _footer_note(doc, ctx):
    from docx.shared import Pt

    p = doc.add_paragraph(
        f"Erstellt am {_de(date.today().isoformat())} · Gewichtung: große Noten "
        f"{ctx['gross_anteil']} %, kleine Noten {100 - ctx['gross_anteil']} % · "
        "Notentendenzen: 1+ = 0,75 · 1 = 1,00 · 1− = 1,25 usw."
    )
    p.runs[0].font.size = Pt(8)


def _class_line(ctx) -> str:
    cls = ctx["class"]
    parts = [f"Klasse {cls['name']}"]
    if cls.get("subject") and cls["subject"] != "kein Fach":
        parts.append(cls["subject"])
    if ctx.get("school_year"):
        parts.append(ctx["school_year"])
    parts.append(TERM_LABEL.get(ctx["term"], ctx["term"]))
    return " · ".join(parts)


def _save(doc) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- 1) Notenmatrix
def build_matrix_docx(ctx) -> bytes:
    """Notenübersicht der Klasse im Querformat: Schüler × Leistungen mit Ø-Spalten."""
    doc = _doc(landscape=True)
    _title(doc, "Notenübersicht", _class_line(ctx))

    items = ctx["items"]
    header = ["Schüler/in"] + [
        f"{it['title']}\n{_de(it['date'])} · {it['kind']} ({SIZE_LABEL.get(it['size'], it['size'])})"
        for it in items
    ] + ["Ø groß", "Ø klein", "Ø gesamt"]

    rows = []
    for s in ctx["students"]:
        summary = ctx["summaries"][s["id"]]
        cells = [_note(ctx["grades"].get((it["id"], s["id"]), {}).get("value")) for it in items]
        rows.append([s["name"]] + cells + [
            _num(summary["avg_gross"]), _num(summary["avg_klein"]), _num(summary["avg_gesamt"]),
        ])

    if not items:
        doc.add_paragraph(f"Im {TERM_LABEL.get(ctx['term'], ctx['term'])} ist noch keine "
                          "Leistung erfasst.")
    else:
        _table(doc, header, rows)
    _footer_note(doc, ctx)
    return _save(doc)


# ---------------------------------------------------------------- 2) Einzelblatt je Schüler
def build_student_sheets_docx(ctx) -> bytes:
    """Eine Seite pro Schüler: alle Noten chronologisch, darunter die Durchschnitte."""
    from docx.shared import Pt

    doc = _doc()
    students = ctx["students"]
    if not students:
        _title(doc, "Notenblätter", _class_line(ctx))
        doc.add_paragraph("Für diese Klasse ist noch keine Schülerliste angelegt.")
        return _save(doc)

    for index, s in enumerate(students):
        if index:
            doc.add_page_break()
        _title(doc, s["name"], _class_line(ctx))

        rows = []
        for it in ctx["items"]:
            grade = ctx["grades"].get((it["id"], s["id"]))
            if grade is None:
                continue
            rows.append([
                _de(it["date"]), it["title"], it["kind"],
                SIZE_LABEL.get(it["size"], it["size"]),
                _note(grade["value"]), grade.get("comment") or "",
            ])
        if rows:
            _table(doc, ["Datum", "Leistung", "Anlass", "Wertigkeit", "Note", "Hinweise"],
                   rows, widths_cm=[2.0, 5.0, 3.0, 2.0, 1.5, 4.5])
        else:
            doc.add_paragraph("In diesem Halbjahr liegt noch keine Note vor.")

        summary = ctx["summaries"][s["id"]]
        doc.add_paragraph("")
        line = doc.add_paragraph(
            f"Ø große Noten: {_num(summary['avg_gross'])}   ·   "
            f"Ø kleine Noten: {_num(summary['avg_klein'])}   ·   "
            f"Ø gesamt: {_num(summary['avg_gesamt'])}"
        )
        line.runs[0].bold = True

        vorschlag = _suggestion(summary)
        if summary["term_grade"] is not None:
            text = f"Zeugnisnote: {_note(summary['term_grade'])} (rechnerischer Vorschlag: {vorschlag})"
            if summary.get("term_grade_comment"):
                text += f" – {summary['term_grade_comment']}"
        else:
            text = f"Rechnerischer Vorschlag für die Zeugnisnote: {vorschlag}"
        doc.add_paragraph(text)

        note = doc.add_paragraph(
            f"Gewichtung: große Noten {ctx['gross_anteil']} %, "
            f"kleine Noten {100 - ctx['gross_anteil']} %."
        )
        note.runs[0].font.size = Pt(8)
    return _save(doc)


# ---------------------------------------------------------------- 3) Auswertung einer Leistung
def build_item_docx(ctx, item) -> bytes:
    """Eine Leistung im Detail: Notenspiegel, Durchschnitt, Liste der Schüler."""
    from docx.shared import Pt

    doc = _doc()
    _title(doc, item["title"],
           f"{_de(item['date'])} · {item['kind']} · "
           f"{SIZE_LABEL.get(item['size'], item['size'])}e Note · {_class_line(ctx)}")
    if item.get("note"):
        doc.add_paragraph(item["note"])

    values = []
    rows = []
    for s in ctx["students"]:
        grade = ctx["grades"].get((item["id"], s["id"]))
        value = grade["value"] if grade else None
        if value is not None:
            values.append(value)
        rows.append([s["name"], _note(value), (grade or {}).get("comment") or ""])

    doc.add_paragraph("")
    spiegel = doc.add_paragraph("Notenspiegel")
    spiegel.runs[0].bold = True

    # Halbwerte (2,5 / 3,5 …) werden KEINEM Balken zugeschlagen, sondern unter der Tabelle
    # gesondert ausgewiesen: welche der beiden Noten es wird, entscheidet der Lehrer je
    # Schüler, das darf eine Klassenstatistik nicht vorwegnehmen.
    grenz = [v for v in values if calc.is_borderline(v)]
    zugeordnet = [v for v in values if not calc.is_borderline(v)]
    # int(v + 0.5) statt round(): Python rundet bei .5 zur GERADEN Zahl (round(2.5) == 2).
    # Halbwerte sind hier zwar schon aussortiert, die Regel bleibt aber die verlässlichere.
    counts = {n: 0 for n in range(1, 7)}
    for v in zugeordnet:
        counts[max(1, min(6, int(v + 0.5)))] += 1
    n_zug = len(zugeordnet)
    _table(
        doc,
        ["Note"] + [str(n) for n in range(1, 7)],
        [["Anzahl"] + [str(counts[n]) for n in range(1, 7)],
         ["Anteil"] + [(f"{counts[n] * 100 / n_zug:.0f} %" if n_zug else "–") for n in range(1, 7)]],
    )
    if grenz:
        grenzzeile = doc.add_paragraph(
            f"Auf der Grenze: {len(grenz)} "
            f"({', '.join(_num(v) for v in sorted(grenz))}) – keinem Balken zugeschlagen, "
            "da hier die Entscheidung zwischen zwei Noten offen ist. "
            f"Die Anteile beziehen sich auf die übrigen {n_zug}."
        )
        grenzzeile.runs[0].font.size = Pt(8)
    stat = doc.add_paragraph(
        f"Bewertet: {len(values)} von {len(ctx['students'])} · "
        f"Durchschnitt: {_num(calc.average(values))}"
    )
    stat.runs[0].font.size = Pt(9)

    doc.add_paragraph("")
    _table(doc, ["Schüler/in", "Note", "Hinweise"], rows, widths_cm=[6.0, 2.0, 9.0])
    return _save(doc)


# ---------------------------------------------------------------- 4) Zeugnisnoten-Liste
def build_term_docx(ctx) -> bytes:
    """Kompakte Liste Schüler → Endnote, mit dem rechnerischen Vorschlag daneben."""
    doc = _doc()
    _title(doc, "Zeugnisnoten", _class_line(ctx))

    rows = []
    for s in ctx["students"]:
        summary = ctx["summaries"][s["id"]]
        gesetzt = summary["term_grade"]
        rows.append([
            s["name"],
            _num(summary["avg_gesamt"]),
            _suggestion(summary),
            _note(gesetzt) if gesetzt is not None else "",
            summary.get("term_grade_comment") or "",
        ])
    _table(doc, ["Schüler/in", "Ø gesamt", "Vorschlag", "Zeugnisnote", "Bemerkung"], rows,
           widths_cm=[5.5, 2.0, 2.0, 2.5, 6.0])
    _footer_note(doc, ctx)
    return _save(doc)
